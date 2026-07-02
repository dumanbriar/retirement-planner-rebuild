"""Auditable Excel workbook: every year, every account, every assumption.

Sheets:
  1. Summary          - headline metrics
  2. Assumptions      - every user input and every modeling rule + source
  3. Accumulation     - per-year per-account flows until retirement
  4. Retirement       - full cash-flow & tax detail for every retired year
  5. Account Detail   - start/contrib/withdraw/convert/growth/end per account-year
  6. Strategies       - Roth-conversion comparison, SS claiming grid,
                        Traditional-vs-Roth contribution split
  7. Sensitivity      - scenario table
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import PlanInput, PlanResult

MONEY = '#,##0'
PCT = '0.00%'

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
WARN_FILL = PatternFill("solid", fgColor="FDE8E8")
TITLE_FONT = Font(bold=True, size=14, color="1E3A5F")
SUB_FONT = Font(size=9, color="666666")


def _sheet_header(ws, row: int, headers: list[str]):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=2)


def _autosize(ws, widths: dict[int, int] | None = None, default: int = 13):
    try:
        max_col = ws.max_column
    except Exception:
        max_col = 20
    for c in range(1, max_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = \
            (widths or {}).get(c, default)


def build_workbook(plan: PlanInput, result: PlanResult) -> bytes:
    wb = Workbook()
    m = result.metrics
    persons = plan.persons

    # ------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Retirement Plan — Audit Workbook"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} — " \
               "deterministic projection; every number on screen is reproducible from " \
               "the sheets in this workbook."
    ws["A2"].font = SUB_FONT
    rows = [
        ("Household", " & ".join(p.name for p in persons)),
        ("Retirement year (household)", m.retirement_year),
        ("Nest egg at retirement (nominal)", m.nest_egg_at_retirement),
        ("Nest egg at retirement (today's $)", m.nest_egg_at_retirement_real),
        ("Plan outcome", "FUNDED through plan end" if m.success
         else f"DEPLETED at age {m.depletion_age}"),
        ("Ending net worth (nominal)", m.ending_net_worth),
        ("Ending net worth (today's $)", m.ending_net_worth_real),
        ("Ending after-tax wealth (today's $)", m.ending_after_tax_real),
        ("Lifetime taxes paid (nominal)", m.lifetime_taxes),
        ("Lifetime taxes paid (today's $)", m.lifetime_taxes_real),
        ("Roth conversion strategy used", m.chosen_conversion_strategy),
        ("Social Security claiming ages",
         ", ".join(f"{p.name}: {a}" for p, a in zip(persons, m.ss_claim_ages))),
        ("Annual spending goal (today's $)", plan.annual_spending),
    ]
    r = 4
    for label, val in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=r, column=2, value=val)
        if isinstance(val, float):
            cell.number_format = MONEY
        r += 1
    if result.warnings:
        r += 1
        ws.cell(row=r, column=1, value="Warnings").font = Font(bold=True)
        for w in result.warnings:
            r += 1
            ws.cell(row=r, column=1, value=w).fill = WARN_FILL
    _autosize(ws, {1: 38, 2: 42})

    # --------------------------------------------------------- Assumptions
    ws = wb.create_sheet("Assumptions")
    ws["A1"] = "Inputs & Modeling Assumptions"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "kind: 'modeled' = statutory rule implemented exactly; " \
               "'estimated' = approximation of a real-world quantity; " \
               "'assumed' = user/plan assumption."
    ws["A2"].font = SUB_FONT
    _sheet_header(ws, 4, ["Item", "Value", "Kind", "Source / method"])
    r = 5
    a = plan.assumptions
    user_inputs = [
        ("Annual retirement spending (today's $)", plan.annual_spending, "assumed", "User input"),
    ]
    for i, p in enumerate(persons):
        user_inputs += [
            (f"{p.name}: current age", p.current_age, "assumed", "User input"),
            (f"{p.name}: retirement age", p.retirement_age, "assumed", "User input"),
            (f"{p.name}: plan-end (death) age", p.death_age, "assumed", "User input"),
            (f"{p.name}: SS monthly benefit at FRA (today's $)", p.ss_monthly_at_fra,
             "assumed", "From SSA statement (PIA)"),
            (f"{p.name}: SS claiming age used", m.ss_claim_ages[i], "modeled",
             "Optimized grid" if a.optimize_ss_claiming else "User input"),
        ]
    for acc in plan.accounts:
        user_inputs.append((
            f"Account: {acc.name} ({acc.type.value}, {persons[acc.owner].name})",
            f"balance {acc.balance:,.0f}; contrib {acc.annual_contribution:,.0f}/yr; "
            f"return {acc.rate() * 100:.2f}%"
            + (f"; basis {acc.cost_basis:,.0f}" if acc.cost_basis is not None else ""),
            "assumed", "User input"))
    for l in plan.liabilities:
        future = (f"; starts age {l.start_age}, down payment {l.down_payment:,.0f}"
                  if l.start_age is not None else "")
        user_inputs.append((f"Liability: {l.name}",
                            f"balance {l.balance:,.0f}; rate {l.interest_rate*100:.2f}%; "
                            f"payment {l.annual_payment:,.0f}/yr{future}",
                            "assumed", "User input"))
    for s in plan.income_streams:
        user_inputs.append((f"Income: {s.name} ({persons[s.owner].name})",
                            f"{s.annual_amount:,.0f}/yr ages {s.start_age}-{s.end_age or 'death'}"
                            f"{' +COLA' if s.cola else ''}{'' if s.taxable else ' (non-taxable)'}",
                            "assumed", "User input"))
    conv = a.roth_conversion_strategy.value
    if conv == "custom":
        conv += f" ({a.custom_conversion_amount:,.0f}/yr today's $)"
    user_inputs += [
        ("Assumption: inflation", f"{a.inflation * 100:g}%", "assumed", "User input"),
        ("Assumption: healthcare inflation", f"{a.healthcare_inflation * 100:g}%",
         "assumed", "User input"),
        ("Assumption: state tax rate (flat)", f"{a.state_tax_rate * 100:g}%",
         "assumed", "User input"),
        ("Assumption: pre-retirement tax rate (flat drag)",
         f"{a.pre_retirement_tax_rate * 100:g}%", "assumed", "User input"),
        ("Assumption: taxable dividend yield", f"{a.taxable_dividend_yield * 100:g}%",
         "assumed", "User input"),
        ("Assumption: heir tax rate", f"{a.heir_tax_rate * 100:g}%", "assumed", "User input"),
        ("Assumption: contributions grow with inflation",
         "yes" if a.contributions_grow_with_inflation else "no", "assumed", "User input"),
        ("Assumption: Roth conversion strategy", conv, "assumed", "User input"),
        ("Assumption: optimize SS claiming ages",
         "yes" if a.optimize_ss_claiming else "no", "assumed", "User input"),
        ("Assumption: suggest Traditional/Roth contribution split",
         "yes" if a.optimize_contribution_split else "no", "assumed", "User input"),
        ("Assumption: ACA benchmark premium ($/mo/person, pre-65)",
         a.aca_benchmark_monthly_per_person, "assumed", "User estimate"),
        ("Assumption: pre-65 out-of-pocket healthcare ($/yr/person)",
         float(a.pre65_oop_annual_per_person), "assumed", "User estimate"),
        ("Assumption: Medicare other costs ($/yr/person, 65+)",
         float(a.medicare_other_annual_per_person), "assumed", "User estimate"),
        ("Assumption: pre-retirement MAGI (IRMAA 2-yr lookback)",
         (f"{a.pre_retirement_magi:,.0f}" if a.pre_retirement_magi is not None
          else "auto (wage MAGI when salaries entered, else 1.5x spending)"),
         "assumed" if a.pre_retirement_magi is not None else "estimated",
         "User input" if a.pre_retirement_magi is not None else "Engine estimate"),
    ]
    for label, val, kind, src in user_inputs:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=val)
        ws.cell(row=r, column=3, value=kind)
        ws.cell(row=r, column=4, value=src)
        r += 1
    r += 1
    for note in result.assumption_notes:
        ws.cell(row=r, column=1, value=note["label"])
        ws.cell(row=r, column=2, value=note["value"])
        ws.cell(row=r, column=3, value=note.get("kind", ""))
        ws.cell(row=r, column=4, value=note["source"])
        r += 1
    _autosize(ws, {1: 44, 2: 36, 3: 11, 4: 80})

    # -------------------------------------------------------- Accumulation
    ws = wb.create_sheet("Accumulation")
    ws["A1"] = "Accumulation Phase — year by year until household retirement"
    ws["A1"].font = TITLE_FONT
    base_year = result.years[0].year if result.years else 2026
    inflation = plan.assumptions.inflation

    def deflator(year: int) -> float:
        """Divide a nominal figure by this to get today's dollars (the same
        deflator the dashboard's Today's-$ toggle applies)."""
        return (1 + inflation) ** (year - base_year)

    age_cols = [f"{p.name} age" for p in persons]
    headers = ["Year"] + age_cols + ["Contributions", "Growth (net of drag)",
               "Tax drag / pre-ret. tax", "Realized gains", "AGI",
               "Taxable income", "Federal tax", "State tax", "Penalties",
               "Effective rate", "Forced RMD", "SS reinvested (gross)",
               "Other income", "Home purchase", "Surplus reinvested",
               "Total assets", "Liabilities",
               "Net worth", "Net worth (today's $)",
               "Deflator (÷ nominal → today's $)", "Notes"]
    _sheet_header(ws, 3, headers)
    r = 4
    for y in result.years:
        if y.phase != "accumulation":
            continue
        contrib = sum(ac.contribution for ac in y.accounts) - y.surplus_reinvested
        growth = sum(ac.growth for ac in y.accounts)
        vals = [y.year] + [y.ages[i] if i < len(y.ages) else None
                           for i in range(len(persons))] + \
            [contrib, growth, y.total_tax, y.realized_gains, y.agi,
             y.taxable_income, y.federal_tax, y.state_tax, y.penalties,
             y.effective_rate, y.rmd_total, y.ss_total, y.other_income,
             y.home_purchase, y.surplus_reinvested,
             y.total_assets, y.total_liabilities, y.net_worth,
             y.net_worth_real, deflator(y.year), "; ".join(y.flags)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c > 1 + len(persons) and isinstance(v, float):
                if headers[c - 1] == "Effective rate":
                    cell.number_format = PCT
                elif headers[c - 1].startswith("Deflator"):
                    cell.number_format = "0.0000"
                else:
                    cell.number_format = MONEY
        r += 1
    _autosize(ws)

    # ---------------------------------------------------------- Retirement
    ws = wb.create_sheet("Retirement")
    ws["A1"] = "Retirement Phase — full cash-flow and tax detail"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("W/D HSA (medical) is tax-free and already netted inside "
                "'Healthcare (net)'; W/D HSA (non-medical) funds spending and is "
                "taxable. The dashboard's 'Wd HSA' column shows the non-medical "
                "figure.")
    ws["A2"].font = SUB_FONT
    headers = (["Year"] + age_cols +
               ["Filing", "Contributions", "Growth (net of drag)",
                "Spend goal", "Debt pmts", "Home purchase", "Healthcare (net)",
                "ACA subsidy", "IRMAA", "SS gross", "Taxable SS", "Other income",
                "RMD", "W/D cash", "W/D taxable", "W/D tax-def", "W/D Roth",
                "W/D HSA (medical)", "W/D HSA (non-medical)", "Roth conversion",
                "Realized gains", "Dividends",
                "Interest", "AGI", "MAGI", "Deductions", "Taxable income",
                "Federal tax", "of which LTCG tax", "NIIT", "State tax",
                "Penalties", "Total tax", "Marginal rate", "Effective rate",
                "Surplus reinvested", "SHORTFALL", "Total assets", "Liabilities",
                "Net worth", "Net worth (today's $)",
                "Deflator (÷ nominal → today's $)", "Notes"])
    _sheet_header(ws, 3, headers)
    r = 4
    for y in result.years:
        if y.phase != "retirement":
            continue
        wbt = y.withdrawals_by_type
        hsa_total = sum(ac.withdrawal for ac in y.accounts if ac.type.value == "hsa")
        hsa_nonmedical = wbt.get("hsa", 0.0)
        hsa_medical = max(0.0, hsa_total - hsa_nonmedical)
        contrib = sum(ac.contribution for ac in y.accounts) - y.surplus_reinvested
        growth = sum(ac.growth for ac in y.accounts)
        vals = ([y.year] + [y.ages[i] if i < len(y.ages) else None
                            for i in range(len(persons))] +
                [y.filing_status, contrib, growth,
                 y.spend_goal, y.debt_payments, y.home_purchase,
                 y.healthcare_cost,
                 y.aca_subsidy, y.irmaa_surcharge, y.ss_total, y.taxable_ss,
                 y.other_income, y.rmd_total, wbt.get("cash", 0.0),
                 wbt.get("taxable", 0.0), wbt.get("tax_deferred", 0.0),
                 wbt.get("roth", 0.0), hsa_medical, hsa_nonmedical,
                 y.roth_conversion,
                 y.realized_gains, y.dividends, y.interest, y.agi, y.magi,
                 y.deductions, y.taxable_income, y.federal_tax, y.ltcg_tax,
                 y.niit, y.state_tax, y.penalties, y.total_tax,
                 y.marginal_rate, y.effective_rate, y.surplus_reinvested,
                 y.shortfall, y.total_assets, y.total_liabilities, y.net_worth,
                 y.net_worth_real, deflator(y.year), "; ".join(y.flags)])
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, float):
                if headers[c - 1] in ("Marginal rate", "Effective rate"):
                    cell.number_format = PCT
                elif headers[c - 1].startswith("Deflator"):
                    cell.number_format = "0.0000"
                else:
                    cell.number_format = MONEY
            if y.shortfall > 1:
                cell.fill = WARN_FILL
        r += 1
    _autosize(ws, default=12)

    # ------------------------------------------------------ Account Detail
    ws = wb.create_sheet("Account Detail")
    ws["A1"] = "Every account, every year: start + contributions - withdrawals " \
               "± conversions + growth = end"
    ws["A1"].font = TITLE_FONT
    _sheet_header(ws, 3, ["Year", "Phase", "Account", "Type", "Owner",
                          "Start balance", "Contribution*", "Withdrawal",
                          "Conversion out", "Conversion in", "Growth", "End balance",
                          "Cost basis (taxable)"])
    ws["A2"] = ("*Contribution includes reinvested amounts, not just new savings: "
                "net pre-retirement inflows (SS/RMD/other income) and reinvested "
                "Traditional tax savings in accumulation years, and reinvested "
                "surplus income in retirement years.")
    ws["A2"].font = SUB_FONT
    r = 4
    for y in result.years:
        for ac in y.accounts:
            vals = [y.year, y.phase, ac.name, ac.type.value,
                    persons[ac.owner].name if ac.owner < len(persons) else "",
                    ac.start_balance, ac.contribution, ac.withdrawal,
                    ac.conversion_out, ac.conversion_in, ac.growth,
                    ac.end_balance, ac.cost_basis]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                if c >= 6 and isinstance(v, float):
                    cell.number_format = MONEY
            r += 1
    _autosize(ws, {3: 28})

    # ---------------------------------------------------------- Strategies
    ws = wb.create_sheet("Strategies")
    ws["A1"] = "Roth Conversion Strategy Comparison (objective: ending after-tax " \
               "wealth, today's $)"
    ws["A1"].font = TITLE_FONT
    _sheet_header(ws, 3, ["Strategy", "Ending after-tax wealth (today's $)",
                          "Lifetime taxes (today's $)", "Total converted",
                          "Depleted at age", "Chosen"])
    r = 4
    for cmp_ in result.conversion_comparison:
        vals = [cmp_.strategy, cmp_.ending_after_tax_real, cmp_.lifetime_taxes_real,
                cmp_.total_converted, cmp_.depletion_age or "",
                "<= chosen" if cmp_.strategy == m.chosen_conversion_strategy else ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, float):
                cell.number_format = MONEY
        r += 1
    if result.ss_grid:
        r += 2
        ws.cell(row=r, column=1, value="Social Security claiming-age grid").font = TITLE_FONT
        r += 1
        _sheet_header(ws, r, [f"{p.name} claims at" for p in persons] +
                      ["Ending after-tax wealth (today's $)", "Depleted at age", "Best"])
        r += 1
        for cell_ in result.ss_grid:
            vals = list(cell_.claim_ages) + [cell_.ending_after_tax_real,
                                             cell_.depletion_age or "",
                                             "<= chosen" if cell_.claim_ages == m.ss_claim_ages else ""]
            for c, v in enumerate(vals, 1):
                xc = ws.cell(row=r, column=c, value=v)
                if isinstance(v, float):
                    xc.number_format = MONEY
            r += 1
    if result.contribution_split:
        r += 2
        ws.cell(row=r, column=1,
                value="Traditional vs. Roth contribution split — advisory "
                      "suggestion (objective: ending after-tax wealth, today's "
                      "$). The projection models the contributions as entered."
                      ).font = TITLE_FONT
        r += 1
        _sheet_header(ws, r, ["Household % to Roth", "Best Roth conversions",
                              "Ending after-tax wealth (today's $)",
                              "Lifetime taxes (today's $)", "Depleted at age",
                              "Current", "Suggested"])
        r += 1
        chosen = round(m.chosen_contribution_split[0], 4) if m.chosen_contribution_split else None
        for cell_ in result.contribution_split:
            vals = [round(cell_.roth_pct[0] * 100), cell_.conversion_strategy,
                    cell_.ending_after_tax_real, cell_.lifetime_taxes_real,
                    cell_.depletion_age or "",
                    "current" if cell_.is_current else "",
                    "<= suggested" if round(cell_.roth_pct[0], 4) == chosen else ""]
            for c, v in enumerate(vals, 1):
                xc = ws.cell(row=r, column=c, value=v)
                if isinstance(v, float):
                    xc.number_format = MONEY
            r += 1
    _autosize(ws, {1: 20, 2: 30, 3: 26, 4: 16})

    # --------------------------------------------------------- Sensitivity
    ws = wb.create_sheet("Sensitivity")
    ws["A1"] = "Sensitivity — how outcomes move when one assumption changes"
    ws["A1"].font = TITLE_FONT
    _sheet_header(ws, 3, ["Scenario", "Parameter", "Change",
                          "Nest egg at retirement (today's $)",
                          "Ending net worth (today's $)", "Depleted at age", "Funded?"])
    r = 4
    base = ["BASE PLAN", "—", "—", m.nest_egg_at_retirement_real,
            m.ending_net_worth_real, m.depletion_age or "",
            "YES" if m.success else "NO"]
    for c, v in enumerate(base, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(bold=True)
        if isinstance(v, float):
            cell.number_format = MONEY
    r += 1
    for srow in result.sensitivity:
        vals = [srow.label, srow.parameter, srow.delta, srow.nest_egg_real,
                srow.ending_net_worth_real, srow.depletion_age or "",
                "YES" if srow.success else "NO"]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, float):
                cell.number_format = MONEY
            if not srow.success:
                cell.fill = WARN_FILL
        r += 1
    _autosize(ws, {1: 30, 4: 30, 5: 28})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
